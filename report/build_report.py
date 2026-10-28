#!/usr/bin/env python3
"""Builds the ECS7032P Assignment 1 report from the IEEE Word template of the module.

    python3 make_figures.py && python3 build_report.py      ->  report.docx and report.pdf

The module template (Report tips and templates/Paper_Template.docx) is stored in strict OOXML, which python-docx
cannot read, so it is first re-saved as a normal .docx with LibreOffice. Text markup used below: <i>..</i>,
<b>..</b>, <sub>..</sub>, <sup>..</sup>; {key} is a citation, numbered in order of first use (IEEE style).
Every number in the text comes from groupA_code/experiments_v3, _v4 and _v5 (RESULTS.md files).
"""
import copy
import os
import re
import subprocess
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "Report tips and templates", "Paper_Template.docx")
WORK = os.path.join(HERE, "build")
OUT = os.path.join(HERE, "ECS7032P_Assignment1_Report_GroupA")

# ----------------------------------------------------------------------------------------------- data
import csv as _csv
import math as _math

ROOT = os.path.normpath(os.path.join(HERE, ".."))
R4 = os.path.join(ROOT, "groupA_code", "experiments_v4", "results")
R5 = os.path.join(ROOT, "groupA_code", "experiments_v5", "results")


def _load(path):
    with open(path) as f:
        return {int(r["game"]): r for r in _csv.DictReader(f)}


def _wilson(p, n, z=1.96):
    c = (p + z * z / (2 * n)) / (1 + z * z / n)
    h = z * _math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return 100 * (c - h), 100 * (c + h)


def _mean_se(xs):
    n = len(xs)
    m = sum(xs) / n
    return m, _math.sqrt(sum((x - m) ** 2 for x in xs) / (n - 1) / n)


def run(path):
    """win %, Wilson CI, score margin and its 1.96 SE of a run (None if the run is still missing)"""
    if not os.path.exists(path):
        return None
    rows = _load(path)
    n = len(rows)
    p = sum(float(r["focusWin"]) for r in rows.values()) / n
    lo, hi = _wilson(p, n)
    m, se = _mean_se([float(r["focusScore"]) - float(r["bestOppScore"]) for r in rows.values()])
    return {"n": n, "win": 100 * p, "lo": lo, "hi": hi, "m": m, "mci": 1.96 * se}


def paired(base, other):
    """paired differences (same seed): win in percentage points and score margin, each with 1.96 SE and z"""
    if not (os.path.exists(base) and os.path.exists(other)):
        return None
    b, o = _load(base), _load(other)
    g = sorted(set(b) & set(o))
    mg = lambda r: float(r["focusScore"]) - float(r["bestOppScore"])
    w, sw = _mean_se([float(o[k]["focusWin"]) - float(b[k]["focusWin"]) for k in g])
    d, sd = _mean_se([mg(o[k]) - mg(b[k]) for k in g])
    return {"dw": 100 * w, "dwci": 196 * sw, "zw": w / sw, "dm": d, "dmci": 1.96 * sd, "zm": d / sd}


def seat_wins(log):
    """win % of every agent's seats from the summary block of an arena log"""
    out = {}
    if os.path.exists(log):
        for line in open(log):
            m = re.match(r"\s+(\S+\.json)(?: \[focus\])?\s+win ([\d.]+)% over", line)
            if m:
                out[m.group(1)] = float(m.group(2))
    return out


def fmt_run(r):
    return "pending" if r is None else f"{r['win']:.1f} [{r['lo']:.1f}, {r['hi']:.1f}]"


def fmt_margin(r):
    return "pending" if r is None else f"{r['m']:+.2f} ± {r['mci']:.2f}".replace("-", "−")


def pp(x, ci, z=None):
    s = f"{x:+.1f} ± {ci:.1f} pp".replace("-", "−")
    return s + (f" (<i>z</i> = {abs(z):.1f})" if z is not None else "")




import glob as _glob
import hashlib as _hashlib


R3 = os.path.join(ROOT, "groupA_code", "experiments_v3", "results")


def seat_wins_v4(label):
    """win % of the opponents' seats in a run of experiments_v3/_v4, from the summary block after its RESULT line"""
    out = {}
    for d in (R3, R4):
        for log in _glob.glob(os.path.join(d, "*.log")):
            lines = open(log, errors="ignore").read().splitlines()
            for i, line in enumerate(lines):
                if line.startswith(f"RESULT {label} "):
                    for x in lines[i + 1:i + 12]:
                        m = re.match(r"\s+(\S+\.json)\s+win ([\d.]+)% over (\d+) seats", x)
                        if m and "[focus]" not in x:
                            out[m.group(1)] = float(m.group(2))
                        if x.startswith("RESULT") or x.startswith("BATCH"):
                            break
    return out


def pooled(paths):
    """win rate (%) of one or several independent runs pooled, and the variance of that mean (%^2)"""
    ws = []
    for f in paths:
        ws += [float(r["focusWin"]) for r in _load(f).values()]
    m, se = _mean_se(ws)
    return 100 * m, (100 * se) ** 2


def two_sample(base, other):
    """difference other - base of (pooled) win rates of independent runs: pp, 1.96 SE, z"""
    mb, vb = pooled(base)
    mo, vo = pooled(other)
    se = _math.sqrt(vb + vo)
    return {"dw": mo - mb, "dwci": 1.96 * se, "zw": (mo - mb) / se}


def diff(a, b):
    """difference b - a of two independent estimates (dicts with dw, dwci)"""
    se = _math.hypot(a["dwci"] / 1.96, b["dwci"] / 1.96)
    return {"dw": b["dw"] - a["dw"], "dwci": 1.96 * se, "zw": (b["dw"] - a["dw"]) / se}


def paired_dd(a1, b1, a2, b2):
    """per-deal difference of differences (b2 - a2) - (b1 - a1) of four runs on the same deals: pp, 1.96 SE, z"""
    runs = [_load(x) for x in (a1, b1, a2, b2)]
    g = sorted(set.intersection(*[set(r) for r in runs]))
    w = lambda r, k: float(r[k]["focusWin"])
    m, se = _mean_se([(w(runs[3], k) - w(runs[2], k)) - (w(runs[1], k) - w(runs[0], k)) for k in g])
    return {"dw": 100 * m, "dwci": 196 * se, "zw": m / se, "n": len(g)}


def study_size():
    """complete runs (at least 200 games; identical copies in several folders counted once; a run joined from two
    parts, X_all.csv, counted instead of its parts), their games and the runs with four players"""
    seen = {}
    for d in ("experiments_v3", "experiments_v4", "experiments_v5"):
        files = _glob.glob(os.path.join(ROOT, "groupA_code", d, "results", "*.csv"))
        joined = {f[:-8] + ".csv" for f in files if f.endswith("_all.csv")}
        for f in files:
            if f in joined or os.path.basename(f) == "r2_field_3p100b.csv":   # parts of a joined run
                continue
            data = open(f, "rb").read()
            with open(f) as fh:
                r = _csv.DictReader(fh)
                if "game" not in (r.fieldnames or []):
                    continue
                rows = list(r)
            if len(rows) >= 200:
                seen[_hashlib.md5(data).hexdigest()] = (len(rows), max(int(x["focusSeat"]) for x in rows) + 1)
    return len(seen), sum(n for n, _ in seen.values()), sum(1 for _, p in seen.values() if p == 4)


def probe_iterations():
    """iterations per decision of v3 and v4 (results/probe_iterations.txt): {(budget, 'v3'): n}"""
    out = {}
    f = os.path.join(R5, "probe_iterations.txt")
    if os.path.exists(f):
        for line in open(f):
            m = re.match(r"(\d+) ms (v\d): .*iterations/decision=(\d+)", line)
            if m:
                out[(int(m.group(1)), m.group(2))] = int(m.group(3))
    return out


def diag4():
    """4-player search diagnostics of the final 500 ms agent (bench4.txt, treestats4.txt, agree4.txt)"""
    try:
        b, t, a = (open(os.path.join(R5, x)).read() for x in ("bench4.txt", "treestats4.txt", "agree4.txt"))
    except OSError:
        return None
    it = dict((k, int(v)) for k, v in re.findall(r"(\S+?\.json)\s+decisions=\d+\s+mean tree iterations=(\d+)", b))
    ts = re.search(r"states (\d+), iterations (\d+) .*root children (\d+), depth-1 visits 1:(\d+)%.*"
                   r"mean depth ([\d.]+)", t)
    seed = re.search(r"decisions (\d+), same card with another seed ([\d.]+)%", a)
    vs = dict((k, float(v)) for k, v in re.findall(r"vs (\S+?\.json)\s+([\d.]+)% same", a))
    if not (it and ts and seed and vs):
        return None
    return {"itF": it.get("arena/agents/final/SushiMCTS_500_FL8.json"), "itNoF": it.get("arena/v5/agents/b500_4noF.json"),
            "states": int(ts.group(1)), "iters": int(ts.group(2)), "kids": int(ts.group(3)), "once": int(ts.group(4)),
            "depth": float(ts.group(5)), "seedSame": float(seed.group(2)), "rerun": vs.get("SushiMCTS_500_FL8.json"),
            "L16": vs.get("b500_4L16.json"), "t1000": vs.get("b1000_4FL8.json")}


def num(x, d=0):
    """number with a thin space as thousands separator"""
    return f"{x:,.{d}f}".replace(",", " ")


R5f = lambda x: os.path.join(R5, x)
R4f = lambda x: os.path.join(R4, x)
first = lambda *xs: next((x for x in xs if os.path.exists(x)), xs[-1])
F3P900 = first(R5f("r2_val_3p900_all.csv"), R5f("r2_val_3p900.csv"))   # 250 + 250 games (merge_parts.py)
FFIELD = first(R5f("v5_field_3p100_all.csv"), R5f("v5_field_3p100.csv"))
VALF = {"3p100": R5f("r2_val_3p100.csv"), "3p500": R5f("r2_val_3p500.csv"), "3p900": F3P900,
        "4p100": R5f("r2_val_4p100.csv"), "4p500": R5f("r2_val_4p500.csv")}
VAL = {k: run(f) for k, f in VALF.items()}
V3H = {k: run(R5f(f"r2_v3_{k}.csv")) for k in ("3p100", "3p500")}               # v3 on the held-out deals
GAIN = {k: paired(R5f(f"r2_v3_{k}.csv"), VALF[k]) for k in ("3p100", "3p500")}  # S + I on the held-out deals
RISE = (paired_dd(R5f("r2_v3_3p100.csv"), VALF["3p100"], R5f("r2_v3_3p500.csv"), VALF["3p500"])
        if all(os.path.exists(x) for x in (R5f("r2_v3_3p100.csv"), R5f("r2_v3_3p500.csv"), VALF["3p500"])) else None)
BUD = {k: paired(VALF["3p100"], VALF[k]) for k in ("3p500", "3p900")}           # final agent, same deals
DUCT = run(R5f("r2_duct_3p100.csv"))
DUCT_SEATS = seat_wins(R5f("r2_duct_3p100.log"))
LEAD = paired(R5f("v5_base3p100.csv"), R5f("r2_lead_3p100.csv"))
OLD = {k: run(R4f(f"V_{k}.csv")) for k in ("3p100", "3p500", "3p900", "4p100", "4p500")}
FIELD = run(FFIELD)
BMR = run(R4f("B_mcts_rhea.csv"))
BOM = run(R4f("B_osla_mcts.csv"))
BSEATS = {**seat_wins_v4("B_mcts_rhea"), **{"osla_" + k: x for k, x in seat_wins_v4("B_osla_mcts").items()}}
BASE5 = run(R5f("v5_base3p100.csv"))
OVF = {"old": two_sample([R4f("V_3p100.csv")], [VALF["3p100"]]),            # held-out vs seed block 7
       "new": two_sample([R5f("v5_base3p100.csv")], [VALF["3p100"]]),
       "machines": paired(R4f("V_3p100.csv"), R5f("v5_base3p100.csv"))}
LAST = {k: paired(R5f("v5_base3p100.csv"), R5f(f"v5_{k}_3p100.csv")) for k in ("X3", "N05", "V2D", "W1", "R3V2")}
SALONE = paired(R4f("A_noSI.csv"), R4f("A_noI.csv"))      # S added to v3 (seed block 7, paired)
VALSEATS = seat_wins(R5f("r2_val_3p100.log"))
# development with three players (seed 1, TAG's own deals: independent samples)
DEVF = {"v2": ["s1_ref", "f5_ref", "f9_ref"], "v3": ["abl_T4", "f5_finalT4", "f9_finalT4"], "v4": ["d1_I", "f5_I", "f9_I"]}
DEV = {k: [run(R4f(x + ".csv")) for x in xs] for k, xs in DEVF.items()}
DEVD = {"v2v3": two_sample([R4f("s1_ref.csv")], [R4f("abl_T4.csv")]),
        "v3drop": two_sample([R4f("abl_T4.csv")], [R4f("f9_finalT4.csv")]),
        "v4drop": two_sample([R4f("d1_I.csv")], [R4f("f9_I.csv")]),
        "si100": two_sample([R4f("abl_T4.csv")], [R4f("d1_I.csv")]),
        "si900": two_sample([R4f("f9_finalT4.csv")], [R4f("f9_I.csv")]),
        "s100": two_sample([R4f("abl_T4.csv")], [R4f("d1_S.csv")]),
        "s500": two_sample([R4f("f5_finalT4.csv")], [R4f("f5_S.csv")])}
DEVD["rise"] = diff(DEVD["si100"], DEVD["si900"])
DEVD["srise"] = diff(DEVD["s100"], DEVD["s500"])
ISM = {k: [seat_wins_v4(x).get("tagmcts.json") for x in xs] for k, xs in DEVF.items()}   # TAG ISMCTS seats
# four players (development; † runs on TAG's own deals are independent samples)
FOUR = {"cfg3": [run(R4f(x + ".csv")) for x in ("q4_I", "q4_Io", "val_q5")],
        "x0": two_sample([R4f("q4_I.csv"), R4f("q4_Io.csv")], [R4f("q4_X0.csv"), R4f("q4_X0c.csv")]),
        "w2": two_sample([R4f("val_q5.csv")], [R4f("q5_W2.csv"), R4f("q5_W2r.csv")]),
        "fl500": paired(R4f("q5_W2r.csv"), R4f("q5_FL8.csv")),
        "fl100": paired(R4f("q4_X0c.csv"), R4f("q4_X0FL8.csv")),
        "cost3": paired(R4f("f5_FL8.csv"), R4f("f5_W2FL8.csv")),
        "depth1": paired(R4f("q5_FL8.csv"), R4f("q5_FL8d1.csv")),
        "b100": paired(R4f("q5_FL8.csv"), R4f("q5_X0b100.csv")),
        "t1000": paired(R4f("q5_FL8.csv"), R4f("q5_FL8t1000K14.csv")),
        "x0_3p": two_sample([R4f("d1_I.csv")], [R4f("d1_X0.csv")])}       # X0 alone with three players, 100 ms
FOUR["inter"] = diff(FOUR["cost3"], FOUR["w2"])
FOUR["x0_inter"] = diff(FOUR["x0_3p"], FOUR["x0"])
ITER = probe_iterations()
DIAG = diag4()
NRUNS, NGAMES, N4RUNS = study_size()


def v(k, field="win", d=1):
    r = VAL[k]
    return "[pending]" if r is None else f"{r[field]:.{d}f}"


def dpp(d, z=True):
    """paired win-rate difference: '−1.3 ± 5.8 pp (z = 0.4)'"""
    return pp(d["dw"], d["dwci"], d["zw"] if z else None)


# ----------------------------------------------------------------------------------------------- content
TITLE = "Decoupled Information-Set MCTS with a Card-Value Model for Sushi Go!"

AUTHORS = [  # three columns, as in the template; fill in the group's names and student numbers
    ["Halit Ozger", "Student no. [to be completed]", "Queen Mary University of London", "London, UK"],
    ["[Group member 2]", "Student no. [to be completed]", "Queen Mary University of London", "London, UK"],
    ["[Group member 3]", "Student no. [to be completed]", "Queen Mary University of London", "London, UK"],
]

FIRST_PAGE_NOTE = ("All sections, tables and figures of this paper were drafted with Claude {claude}; the prompts are "
                   "given in the statement on the use of generative AI at the end of the paper.")

def _abstract_end():
    if RISE is not None and RISE["zw"] >= 1.96:
        return "the decoupled information-set tree gains more over our previous agent at 500 than at 100 ms."
    return (f"settings tuned for four players cost {-FOUR['cost3']['dw']:.1f} percentage points with three "
            "(exploratory).")


ABSTRACT = (
    "In <i>Sushi Go!</i>, players pick cards simultaneously from hidden hands that are passed on, which basic Monte "
    "Carlo tree search (MCTS) does not model. To BasicMCTS of the Tabletop Games (TAG) framework we added rollouts "
    "guided by a card-value model, a lead-based evaluation, round-end rollouts, max-n backups, decoupled selection and "
    "information-set search. On held-out deals the agent wins "
    f"{v('3p100')}%, {v('3p500')}% and {v('3p900')}% of 500 three-player games at 100, 500 and 900 ms against TAG's "
    f"information-set MCTS and our previous agent (chance 33.3%), and {v('4p100')}% of four-player games at 100 ms "
    "(chance 25%). In leave-one-out ablations every component improves the score margin significantly, the "
    "evaluation and the decoupled information-set tree also the win rate; " + _abstract_end())

KEYWORDS = "Monte Carlo tree search, information sets, simultaneous moves, card games, Sushi Go!, TAG"

REFS = {
    "tag": "R. D. Gaina, M. Balla, A. Dockhorn, R. Montoliu, and D. Perez-Liebana, “TAG: A tabletop games "
           "framework,” in <i>Proc. AIIDE Workshop on Experimental AI in Games (EXAG)</i>, 2020.",
    "browne": "C. B. Browne <i>et al.</i>, “A survey of Monte Carlo tree search methods,” <i>IEEE Trans. Comput. "
              "Intell. AI Games</i>, vol. 4, no. 1, pp. 1–43, 2012.",
    "cowling": "P. I. Cowling, E. J. Powley, and D. Whitehouse, “Information set Monte Carlo tree search,” "
               "<i>IEEE Trans. Comput. Intell. AI Games</i>, vol. 4, no. 2, pp. 120–143, 2012.",
    "klang": "C.-M. E. Klang, V. Enhörning, A. Alvarez, and J. Font, “Assessing simultaneous action selection and "
             "complete information in TAG with Sushi Go!,” in <i>Proc. IEEE Conf. Games (CoG)</i>, 2021, pp. 1–4.",
    "kocsis": "L. Kocsis and C. Szepesvári, “Bandit based Monte-Carlo planning,” in <i>Proc. 17th Eur. Conf. "
              "Mach. Learn. (ECML)</i>, 2006, pp. 282–293.",
    "silver": "D. Silver and G. Tesauro, “Monte-Carlo simulation balancing,” in <i>Proc. 26th Int. Conf. Mach. "
              "Learn. (ICML)</i>, 2009, pp. 945–952.",
    "robilliard": "D. Robilliard, C. Fonlupt, and F. Teytaud, “Monte-Carlo tree search for the game of ‘7 Wonders’,” in "
                  "<i>Computer Games (CGW 2014)</i>, Commun. Comput. Inf. Sci., vol. 504. Cham, Switzerland: Springer, "
                  "2014, pp. 64–77.",
    "long": "J. Long, N. R. Sturtevant, M. Buro, and T. Furtak, “Understanding the success of perfect information "
            "Monte Carlo sampling in game tree search,” in <i>Proc. 24th AAAI Conf. Artif. Intell.</i>, 2010, "
            "pp. 134–140.",
    "frank": "I. Frank and D. Basin, “Search in games with incomplete information: A case study using Bridge card "
             "play,” <i>Artif. Intell.</i>, vol. 100, no. 1–2, pp. 87–123, 1998.",
    "lanctot": "M. Lanctot, V. Lisý, and M. H. M. Winands, “Monte Carlo tree search in simultaneous move games "
               "with applications to Goofspiel,” in <i>Computer Games (CGW 2013)</i>, Commun. Comput. Inf. Sci., "
               "vol. 408. Cham, Switzerland: Springer, 2014, pp. 28–43.",
    "tak": "M. J. W. Tak, M. Lanctot, and M. H. M. Winands, “Monte Carlo tree search variants for simultaneous "
           "move games,” in <i>Proc. IEEE Conf. Comput. Intell. Games (CIG)</i>, 2014, pp. 232–239.",
    "sturtevant": "N. R. Sturtevant, “An analysis of UCT in multi-player games,” in <i>Computers and Games "
                  "(CG 2008)</i>, Lect. Notes Comput. Sci., vol. 5131. Berlin, Germany: Springer, 2008, pp. 37–49.",
    "goodman": "J. Goodman, D. Perez-Liebana, and S. Lucas, “MultiTree MCTS in tabletop games,” in <i>Proc. IEEE "
               "Conf. Games (CoG)</i>, 2022, pp. 292–299.",
    "wilson": "E. B. Wilson, “Probable inference, the law of succession, and statistical inference,” <i>J. Amer. "
              "Statist. Assoc.</i>, vol. 22, no. 158, pp. 209–212, 1927.",
    "claude": "Anthropic, “Claude” [large language model, used through the Claude Code tool], Oct. 2026. [Online]. "
              "Available: https://claude.ai (accessed Oct. 28, 2026).",
}


def fpp(d, z=True, sign=True):
    """'+3.1 ± 6.2 pp (z = 1.0)' from a dict with dw, dwci, zw"""
    s = (f"{d['dw']:+.1f}" if sign else f"{abs(d['dw']):.1f}") + f" ± {d['dwci']:.1f} pp"
    if z:
        s += f" (<i>z</i> = {abs(d['zw']):.1f})"
    return s.replace("-", "−")


def dpm(d):
    """table cell '+4.3 ± 4.3'"""
    return f"{d['dw']:+.1f} ± {d['dwci']:.1f}".replace("-", "−")


def fwin(r):
    return "[pending]" if r is None else f"{r['win']:.1f}%"


B = []
h1 = lambda t: B.append(("h1", t))
h2 = lambda t: B.append(("h2", t))
p = lambda t: B.append(("p", t))
bullet = lambda t: B.append(("bullet", t))

# ------------------------------------------------------------------ I
h1("Introduction")
p("Tabletop games combine chance, hidden information and several interacting players within small rule sets; "
  "the Tabletop Games framework (TAG) {tag} gives many of them a common interface so that agents can be compared. "
  "Monte Carlo tree search (MCTS) {browne} is the usual starting point for such agents, but its basic form assumes "
  "that the state is fully observed and that players move one after another {cowling}. <i>Sushi Go!</i>, one of "
  "TAG's games, breaks both assumptions: every player secretly picks one card from its hand, all players reveal "
  "their cards together and pass the rest of their hands on, and after three rounds the highest score wins. Klang "
  "et al. {klang} introduced the game to TAG; the current version shows an agent only the cards it has seen.")
p("Four properties make the game hard for a search agent. (i) The opponents' hands are hidden until they are "
  "passed on. (ii) Moves are simultaneous: no card can be chosen in reaction to an opponent's card of the same "
  "turn. (iii) Many points are delayed and conditional – sets pay out only when completed, Maki rolls at the end of "
  "the round and Pudding at the end of the game – so short simulations miss them and long ones are noisy. (iv) "
  "Hands of nine (three players) or eight cards (four players) give 729 or 4 096 joint moves per turn (more with "
  "Chopsticks), with 0.1–1 s per decision. TAG's BasicMCTS addresses none of the first three: it searches the one "
  "shuffled sample of the hidden cards it receives as if it were the true state, treats the players of a turn as "
  "moving one after another and evaluates random rollouts by the player's own score.")
p("We ask which adaptations of MCTS to these properties make an agent stronger (RQ1), and how their value changes "
  "with the time per decision (RQ2) and the number of players (RQ3). We extend BasicMCTS in steps, keeping every "
  "modification as a switch, and test each in experiments of 500 games per condition (250 in some runs at 900 ms "
  "and in four-player development at 500 ms). Our contributions are:")
bullet("a model of Sushi Go!'s scoring that estimates every card's marginal value, used as the rollout policy of "
       "all players, and every player's final score, used in a bounded, lead-based evaluation;")
bullet("its combination on BasicMCTS with round-end rollouts, max-n backups, decoupled UCT and information-set "
       "search, compared with TAG's own decoupled information-set MCTS;")
bullet("exploratory evidence that settings tuned for four players hurt with three, and a diagnosis of why time "
       "beyond 500 ms helps little with four (a wide, shallow joint-move tree);")
bullet("leave-one-out and factorial ablations, held-out validation and negative results that locate the agent's "
       "remaining limits.")

# ------------------------------------------------------------------ II
h1("Background")
h2("TAG and Sushi Go!")
p("TAG {tag} splits a game into a state, a forward model and parameters; an agent receives a copy of the state, in "
  "which every card it has not seen is reshuffled, and uses the forward model to simulate moves. In TAG's Sushi Go! "
  "the 108-card deck holds Tempura (5 points per pair), Sashimi (10 per triple), Dumplings (1, 3, 6, 10 or 15 "
  "points for one to five), Salmon, Squid and Egg Nigiri (2, 3 and 1 points, tripled on an earlier Wasabi), Maki "
  "rolls with one to three icons (6 and 3 points for the most and second-most icons of a round), Chopsticks (two "
  "cards in a later turn) and Pudding (+6 for the most and, with three or more players, −6 for the fewest at the "
  "end of the game, when Puddings also break ties). Each round deals 12 − <i>n</i> cards to each of <i>n</i> players. As the hands rotate, every "
  "hand of a round has been seen after <i>n</i> − 1 passes, so hidden information is concentrated at the start of "
  "each round, while the value of a card depends on what the others will take from the hands we pass on.")
h2("Monte Carlo Tree Search")
p("MCTS {browne} builds a search tree by repeating selection, expansion, simulation (rollout) and "
  "backpropagation. Upper Confidence bounds applied to Trees (UCT) {kocsis} selects the child that maximises the "
  "upper confidence bound (UCB1) <i>Q</i> + <i>K</i>√(ln <i>N</i> / <i>n</i>), where <i>Q</i> is the child's mean "
  "reward, <i>n</i> its visits and <i>N</i> those of its parent. Domain knowledge can enter through heuristic "
  "rollout policies or priors such as progressive bias {browne}, but a stronger rollout policy does not guarantee "
  "better estimates: Silver and Tesauro {silver} showed that rollouts should be balanced, i.e., give unbiased value "
  "estimates.")
h2("Hidden Information, Simultaneous Moves and Several Players")
p("Determinisation, or perfect-information Monte Carlo (PIMC) sampling, samples the hidden information, searches "
  "each sample as if it were the true state and combines the results. Long et al. {long} related its success in "
  "card games to properties such as how quickly hidden information is revealed during play, which in Sushi Go! "
  "happens within <i>n</i> − 1 turns. Its known flaws are strategy fusion – the search assumes that it can act "
  "differently in states it cannot tell apart – and non-locality {frank}. Information-set MCTS (ISMCTS) {cowling} "
  "avoids strategy fusion by searching one tree over the searcher's information sets and sampling the hidden "
  "information anew in every iteration. In its single-observer form the opponents' choices in the tree are not "
  "conditioned on their own private information; the multiple-observer form models it with one tree per player at "
  "a higher cost {cowling}, an idea TAG's advanced MCTS agent offers as MultiTree MCTS {goodman}.")
p("Simultaneous moves are usually handled by decoupled UCT (DUCT): every player chooses its own action at a node "
  "from its own bandit statistics and the joint action leads to a child {lanctot}. UCB-based DUCT need not converge "
  "to an equilibrium, so Lanctot et al. {lanctot} compared it in Goofspiel with EXP3, regret matching and their "
  "online outcome sampling: regret matching and outcome sampling played best, and decoupled UCT's strategies were "
  "the most exploitable. Tak et al. {tak} found across nine games that the best variant depends on the game and the "
  "opponent, although decoupled UCT was best overall. With several players, UCT with max-n backups "
  "(every player maximises its own value) computes a mixed-strategy equilibrium and played at least as well as "
  "earlier algorithms {sturtevant}.")
p("MCTS has also been applied to the drafting game 7 Wonders, which passes hands in the same way, and many results "
  "known from classic games held there too {robilliard}. "
  "TAG's advanced MCTS agent already offers information sets, one tree per player and a decoupled option; we build "
  "on BasicMCTS, which has none of them, so that the contribution of each idea can be measured. We know of no "
  "study that separates the contributions of decoupled selection, information sets and domain knowledge within one "
  "agent or measures how they change with the time budget and the number of players; the earlier Sushi Go! study "
  "in TAG {klang} focused on "
  "simultaneous action selection and complete information and concluded that better heuristics should improve the "
  "agents, which motivates our card-value model.")

# ------------------------------------------------------------------ III
h1("Method")
p("Our agent, SushiMCTS, is a modified copy of TAG's BasicMCTS (package <i>players.basicMCTS</i>) and uses no code "
  "of TAG's advanced MCTS agent. Every modification is a switch in its parameter class (Table I). We compare the "
  "final agent (v4) with two earlier versions, v2 (H, A, B, D) and v3 (H, B, D, P, R, M), whose tree treats the "
  "players of a turn as moving in sequence (FIGREF_TURN). The search techniques come from the literature; our own "
  "element is the card-value model shared by the evaluation and every player's simulation.")
h2("Card-Value Model and Heuristic (H)")
p("For a state <i>s</i> and player <i>p</i> the model estimates the final score")
B.append(("eq", "Ŝ<sub><i>p</i></sub> = <i>b<sub>p</sub></i> + α<i>u<sub>p</sub></i> + <i>m<sub>p</sub></i> + "
                "π<i>w<sub>t</sub> d<sub>p</sub></i>,", 1))
p("where <i>b<sub>p</sub></i> is the banked score (TAG scores Tempura, Sashimi, Dumplings and Nigiri when they are "
  "revealed); <i>u<sub>p</sub></i> the value of unfinished sets, of unused Wasabi (2ν, the bonus for a Nigiri of "
  "average value ν = 2) and of Chopsticks (1), each weighted by min(1, (<i>r</i> − <i>k</i> + 1)/<i>T</i>), or 0 "
  "if <i>r</i> < <i>k</i>, when <i>k</i> more cards are needed and <i>r</i> picks are left; <i>m<sub>p</sub></i> "
  "the Maki points <i>p</i> would get if the round ended now; and <i>d<sub>p</sub></i> the Pudding (dessert) points "
  "if the game ended now, weighted by <i>w<sub>t</sub></i> = <i>t</i>/3 in round <i>t</i> (α = 0.5, <i>T</i> = 4, "
  "π = 1). The value of <i>s</i> for player <i>i</i> is")
B.append(("eq", "<i>H<sub>i</sub></i>(<i>s</i>) = tanh(Δ<sub><i>i</i></sub> / σ),", 2))
B.append(("eq", "Δ<sub><i>i</i></sub> = <i>w</i>(Ŝ<sub><i>i</i></sub> − max<sub><i>j</i>≠<i>i</i></sub> "
                "Ŝ<sub><i>j</i></sub>) + (1 − <i>w</i>)(Ŝ<sub><i>i</i></sub> − mean<sub><i>j</i>≠<i>i</i></sub> "
                "Ŝ<sub><i>j</i></sub>),", 3))
p("with σ = 15 and <i>w</i> = 0.7: the first term rewards beating the leader, since only first place wins, and "
  "the second keeps a gradient when the agent is far ahead or behind. The marginal value "
  "<i>mv<sub>p</sub></i>(<i>c</i>) of a card <i>c</i> is the change of Ŝ<sub><i>p</i></sub> when <i>p</i> plays it.")
B.append(("table", "Modifications of BasicMCTS (✓ = used in the final agent with 3 / 4 players)",
          [0.5, 2.12, 0.76],
          ["Id", "Modification", "3 p / 4 p"],
          [["H", "heuristic (2)–(3) from model (1)", "✓ / ✓"],
           ["B", "ε-greedy card-value rollouts for every player", "✓ / ✓"],
           ["R", "rollouts end when the round is scored", "✓ / ✓"],
           ["P", "one branch per distinct card", "✓ / ✓"],
           ["M", "max-n backups {sturtevant}", "✓ / ✓"],
           ["S", "decoupled UCT, one node per turn {lanctot}", "✓ / ✓"],
           ["I", "single-observer information-set search {cowling}", "✓ / ✓"],
           ["X0", "four-player card constants <sup>a</sup>", "– / ✓"],
           ["W2", "win-gap value for 4 players (Section III-D)", "– / ✓ <sup>b</sup>"],
           ["F, L", "fast simulator, 8 rollouts per leaf", "– / ✓ <sup>b</sup>"],
           ["A", "progressive bias {browne}", "dropped in v3"],
           ["D", "ensemble of up to 4 determinised trees {cowling}", "replaced by I"]],
          "<sup>a</sup> (σ, <i>w</i>, α, <i>T</i>, β, <i>T</i><sub>m</sub>, κ, δ, π, ν, γ) = (15, 0.7, 0.5, 6, 0.5, 3, "
          "0.25, 1.5, 1.25, 3, 3), default (15, 0.7, 0.5, 4, 1, 0, 0, 0, 1, 2, 0). β scales a Sashimi triple still "
          "needing two cards; the Maki standing has weight 1/(1 + <i>r</i>/<i>T</i><sub>m</sub>) (1 if "
          "<i>T</i><sub>m</sub> = 0), κ per icon the rest; δ weights the next Dumpling; γ: cards needed per "
          "missing card. <sup>b</sup> 500 ms or more."))
B.append(("fig", "figures/fig1_turnnode.png", 3.39, "FIG1_CAPTION"))
h2("Simulation (B, R)")
p("<b>B:</b> every player plays its card with the highest <i>mv</i>, or a random card with probability ε = 0.2 "
  "(the second card of a Chopsticks move is chosen in the same way); the same policy models the opponents. "
  "<b>R:</b> a rollout stops when the current round has been scored; only Pudding carries over, which (1) captures, "
  "so at the leaves <i>u<sub>p</sub></i> = <i>m<sub>p</sub></i> = 0 and only banked points and the Pudding standing "
  "count.")
h2("Tree (P, M, S, I)")
p("<b>P:</b> identical cards lead to identical states, so a node has one branch per distinct card. <b>M:</b> a "
  "node stores one value per player and every player maximises its own value {sturtevant}, instead of BasicMCTS's "
  "paranoid assumption that every opponent minimises ours. <b>S:</b> one node represents a whole turn; at a node "
  "every player <i>p</i> picks a card by")
B.append(("eq", "<i>a<sub>p</sub></i> = argmax<sub><i>a</i></sub> [<i>Q<sub>p</sub></i>(<i>a</i>) + "
                "<i>K</i>√(ln <i>N<sub>p</sub></i>(<i>a</i>) / <i>n<sub>p</sub></i>(<i>a</i>))]", 4))
p("from its own statistics, trying untried cards first; the joint "
  "choice is played in the forward model and leads to the child of that joint action {lanctot}; in v3's tree a "
  "later player “saw” the card an earlier player had just chosen (FIGREF_TURN). <b>I:</b> every iteration starts "
  "from a new sample of the cards we cannot see and walks one shared open-loop tree, whose nodes store statistics "
  "but no states {cowling}. Statistics are keyed by card type; <i>n<sub>p</sub></i>(<i>a</i>) counts how often "
  "<i>p</i> chose <i>a</i> at the node and the availability count <i>N<sub>p</sub></i>(<i>a</i>) how many "
  "iterations reached the node with <i>a</i> in <i>p</i>'s sampled hand. I replaces the determinisation ensemble D "
  "of v2 and v3, which split the budget between up to four independent trees. With three players <i>K</i> = 0.5.")
h2("Four Players")
p("With four players the agent switches parameters. It uses the card constants X0 (Table I), found by a "
  "search-free tuner of greedy card-value policies. From 500 ms it also uses "
  "<b>W2</b> – <i>w</i> = 1; finished games valued by tanh((<i>b<sub>i</sub></i> − max<sub><i>j</i>≠<i>i</i></sub> "
  "<i>b<sub>j</sub></i>)/2) with Pudding as tie-break; σ multiplied by (ρ/2)<sup>0.5</sup> for ρ rounds left; opponents in the backup valuing their lead over the others' mean – and <b>F</b>, TAG's rules "
  "re-implemented on integer arrays and checked move by move against the forward model, whose speed is spent on "
  "<b>L</b> = 8 rollouts per leaf, with <i>K</i> = 1.")
h2("Experimental Setup")
B.append(("table", "Experimental setup", [0.72, 2.66],
          ["Item", "Setting"],
          [["Framework", "TAG as provided for the module, Sushi Go! with its default rules; OpenJDK 21"],
           ["Opponents",
            "v2: our previous agent (H, A, B, D, <i>K</i> 1.41). TAG ISMCTS: SushiGo_3+P_HiBudget at our budget (one "
            "tree per player {goodman}, UCB-Tuned, <i>K</i> 1, random rollouts, score heuristic). TAG DUCT: "
            "SushiGoDUCT (decoupled UCT, <i>K</i> 1, information sets, lead heuristic, 30-move rollouts). BasicMCTS (<i>K</i> 1.41, 10-move rollouts), Rolling Horizon "
            "Evolutionary Algorithm (RHEA; horizon 10, population 10), One-Step Look-Ahead (OSLA, score heuristic)"],
           ["Budget", "CPU time per decision, equal for all agents: 100 ms (most development), 500 and 900 ms; "
                      "submitted agent files: 900 ms (1 s limit) or 430 ms (500 ms limit), with wall-clock guards"],
           ["Seating", "our agent in seat <i>g</i> mod <i>n</i> in game <i>g</i>; the other seats alternate between "
                       "the two opponents (4 players: 2 + 1, swapped every 4 games)"],
           ["Deals", "game <i>g</i> uses seed <i>s</i> + <i>g</i>; from 19 October (all seed-7 and held-out runs) a "
                     "fixed deck order is shuffled with it (equal <i>s</i>, equal deals: paired runs); earlier runs "
                     "used TAG's own deal (independent samples)"],
           ["Seeds", "development <i>s</i> = 1, 2 (v3 also 1001, 2001, 3001); ablation and later candidate tests "
                     "<i>s</i> = 7; held-out <i>s</i> = 100 000, disjoint from all other deals"],
           ["Games", "500 per condition; 250 at 900 ms (development, <i>s</i> = 7) and in four-player development "
                     "at 500 ms (45–70 s of CPU time per game)"],
           ["Hardware, records", "18-core laptop (6–16 games in parallel); candidate tests and validation on a "
                                 "4-core machine (4 in parallel). Agent files, per-game results and scripts: from "
                                 "the authors"]],
          None))
p("Table II summarises the setup. Each game is an independent trial whose outcome for our agent is 1 for a win, "
  "1/<i>k</i> for a <i>k</i>-way shared first place and 0 otherwise. Win rates are reported with Wilson 95% "
  "confidence intervals (CI) {wilson}; the score margin (our score minus the best opponent's) as mean ± 1.96 "
  "standard errors (SE); differences between runs as the mean of the per-game differences on the same deals "
  "± 1.96 SE (for independent samples, the difference of the means with the two-sample SE) and <i>z</i> = "
  "difference/SE, in percentage points (pp) for win rates. With 500 games a win rate near 50% has a 95% interval "
  "of ±4.4 pp, and a difference between two runs about ±5.7 pp even when paired, because the outcomes of two agents "
  "on the same deal correlate only weakly (0.0–0.2); such a comparison detects a difference of about 8 pp with 80% "
  "power. The score margin, usually the more sensitive measure (<i>z</i> values 1.0–2.3 times as large in the "
  "ablation), served to screen ideas. We report unadjusted intervals, apply a Holm correction within the "
  "ablation and within the later candidate tests, and treat the four-player comparisons (" + str(N4RUNS) + " of our "
  + str(NRUNS) + " runs) as exploratory. We expected (H1) every component to improve both measures, (H2) the "
  "decoupled information-set tree to matter more at longer budgets and (H3) the best settings to depend on the "
  "number of players; H2 arose from the development runs and is tested on the held-out deals.")

# ------------------------------------------------------------------ IV
h1("Results")
h2("Final Agent (RQ1)")
B.append(("table", "Final agent on held-out deals and against other agents",
          [0.42, 1.06, 0.28, 0.94, 0.69],
          ["Setting", "Opponents", "<i>n</i>", "Win % [95% CI]", "Margin"],
          [["3p/100", "v2, ISMCTS", str(VAL['3p100']['n']) if VAL['3p100'] else "500", fmt_run(VAL['3p100']), fmt_margin(VAL['3p100'])],
           ["3p/500", "v2, ISMCTS", str(VAL['3p500']['n']) if VAL['3p500'] else "500", fmt_run(VAL['3p500']), fmt_margin(VAL['3p500'])],
           ["3p/900", "v2, ISMCTS", str(VAL['3p900']['n']) if VAL['3p900'] else "500", fmt_run(VAL['3p900']), fmt_margin(VAL['3p900'])],
           ["4p/100", "v2, ISMCTS", str(VAL['4p100']['n']) if VAL['4p100'] else "500", fmt_run(VAL['4p100']), fmt_margin(VAL['4p100'])],
           ["4p/500", "v2, ISMCTS", str(VAL['4p500']['n']) if VAL['4p500'] else "500", fmt_run(VAL['4p500']), fmt_margin(VAL['4p500'])],
           ["3p/100", "DUCT, ISMCTS", str(DUCT['n']) if DUCT else "500", fmt_run(DUCT), fmt_margin(DUCT)],
           ["3p/100", "ISMCTS × 2 <sup>a</sup>", str(FIELD['n']), fmt_run(FIELD), fmt_margin(FIELD)],
           ["3p/100", "BasicMCTS, RHEA <sup>a</sup>", str(BMR['n']), fmt_run(BMR), fmt_margin(BMR)],
           ["3p/100", "OSLA, BasicMCTS <sup>a</sup>", str(BOM['n']), fmt_run(BOM), fmt_margin(BOM)]],
          "Setting: players / ms per decision. Deals: held-out seed block 100 000; <sup>a</sup> seed block 7. Margin: our "
          "score minus the best opponent's in points, ± 95% CI. Chance: 33.3% (3 players), 25% (4 players)."))
p("FINAL_AGENT_PARAGRAPH")
h2("Ablation (RQ1)")
B.append(("fig", "figures/fig3_effects.png", 3.39, "FIG_EFFECTS_CAPTION"))
p("ABLATION_PARAGRAPH")
p("LAST_STUDY_PARAGRAPH")
h2("Time Budget (RQ2)")
B.append(("fig", "figures/fig2_budget.png", 3.39, "FIG_BUDGET_CAPTION"))
p("BUDGET_PARAGRAPH")
h2("Four Players and the Limits of Search (RQ3)")
p("FOUR_PARAGRAPH")

# ------------------------------------------------------------------ V
h1("Conclusions and Future Work")
p("CONCLUSION_1")
p("CONCLUSION_2")

B.append(("h5", "Use of Generative AI"))
B.append(("pai", "We used Claude {claude} (Anthropic), through the Claude Code tool, throughout the project: to write "
          "and modify the Java code of our agent on top of BasicMCTS, to write and run the experiment and statistics "
          "scripts, and to draft all sections, tables and figures of this paper from the logs of our runs; every "
          "number is computed from the result files of our runs. The authors take full responsibility for the code "
          "and the paper, including any errors. The paper was generated with the prompt (original Turkish, then "
          "English translation): <i>“simdi reponun icindeki butun herseyi incele ve hocaya direk yuklyecegimiz "
          "raporuda olusutr repoda ornekler tipler brief hersey var 100 alcagimiz bir rapor olsun her raporu "
          "tamaladiginda baska bir agent spawn edip bagimsiz durusstce kontrol taki yuzde yuz alacagana emin olankadar "
          "looptan cikma”</i> – “Now examine everything in the repository and also create the report we will upload "
          "directly to the lecturer; the repository has examples, tips, the brief, everything. It should be a report "
          "that gets 100. Every time you finish the report, spawn another agent to check it independently and "
          "impartially, and do not leave the loop until you are sure it will get 100%.” Following that prompt, each "
          "draft was marked by a second, independent Claude instance, prompted as “an independent, strict marker for "
          "the QMUL module ECS7032P” to mark it against the brief and the rubric, to recompute its numbers from our "
          "result files and to list the changes that would raise the mark. Its written reports, one per draft (available "
          "on request), were the prompts for the revisions; they asked, among other things, for validation on deals "
          "disjoint from development, a correct description of the availability-count UCB, all replicate runs, "
          "intervals and hedges for non-significant trends, the comparison with TAG's DUCT agent, the lead-heuristic "
          "ablation, v3 on the held-out deals, the four-player diagnostics, correct figure references, this statement's "
          "citation and prompts, claims no stronger than the statistics and the selection bias of v3's development "
          "run. The later candidate changes were tested with the prompt “The lecturer "
          "will hold a tournament of our agents in class; as before, improve the agents' performance quickly while "
          "following the lecturer's restrictions and recording what we did in the README; you have 4.5 hours, test "
          "everything in parallel.” The prompts of earlier development sessions are not reproduced here; they asked "
          "Claude to implement and test the modifications of Table I and to run the experiments of Section IV."))
B.append(("h5", "References"))
B.append(("refs",))


# ----------------------------------------------------------------------------------------------- dynamic text
def count73(pct_same):
    """positions (of DIAG['states']) on which two searches chose different cards"""
    return round(DIAG["states"] * (100 - pct_same) / 100)


def text_fig1_caption():
    return ("The turn node of the final agent (right) and the sequential tree of v3 (left). The final agent re-samples "
            "the cards it cannot see in every iteration (I), lets every player pick by UCB on its own card statistics "
            "(S, P), plays the joint pick, rolls out with card-value policies until the round is scored (B, R) and "
            "backs every player's value (H) up to the card it chose (M); after the time budget it plays its most "
            "visited root card.")


def text_final_agent():
    if VAL["3p100"] is None:
        return "[results of the held-out runs pending]"
    w3 = [VAL[k]["win"] for k in ("3p100", "3p500", "3p900") if VAL[k]]
    w4 = [VAL[k]["win"] for k in ("4p100", "4p500") if VAL[k]]
    rng = lambda w: f"{w[0]:.1f}%" if len(w) == 1 else f"{min(w):.1f}–{max(w):.1f}%"
    s = (f"Table III reports the final agent: on the held-out deals it wins {rng(w3)} of three-player and {rng(w4)} "
         f"of four-player games against v2 and TAG's ISMCTS. On seed block 7, whose deals overlap the development "
         f"deals, it had won {OLD['3p100']['win']:.1f}% at 100 ms ({BASE5['win']:.1f}% when re-run on the machine of "
         f"the held-out runs), {fpp(OVF['old'], z=False, sign=False)} and {fpp(OVF['new'], z=False, sign=False)} more "
         "than on the held-out deals, neither significant.")
    dd = DUCT_SEATS.get("b100_tagduct.json")
    if DUCT and dd is not None:
        s += (f" TAG's DUCT agent, with decoupled selection and information sets but neither our card-value model nor "
              f"round-end rollouts, wins {dd:.1f}% of its seats when it replaces v2"
              + (", the same share as v2." if abs(dd - VALSEATS.get("ref.json", -9)) < 0.05 else "."))
    m = max(BSEATS.get("fw_BasicMCTS.json", 0), BSEATS.get("fw_RHEA.json", 0), BSEATS.get("osla_fw_BasicMCTS.json", 0))
    s += (f" BasicMCTS, RHEA and OSLA win at most {m:.1f}% of their seats; OSLA none, as its one-step look-ahead "
          "never reaches the simultaneous reveal.")
    return s


def text_ablation():
    s = ("FIGREF_EFFECTS (top) removes one component at a time from the final agent. Every removal lowers both "
         "measures; on the score margin all five differences survive a Holm correction, on the win rate only those of "
         "H and of S + I (−R: <i>z</i> = 2.0, −B: 1.4, −I: 1.3), so H1 holds for the margin and only partly for wins. "
         "Without H the leaves are valued by TAG's default heuristic (own score / 50); as <i>K</i> was not retuned "
         "to its scale, the 28.3 ± 5.8 pp loss mixes the value of H with mis-scaled exploration.")
    if LEAD:
        s += (f" TAG's lead heuristic (own score minus the best opponent's) with <i>K</i> rescaled to 7.5 costs "
              f"{fpp(LEAD, sign=False)}: valuing the lead recovers roughly half of H's effect (two estimates against "
              "different base runs). As round-end leaves hold no unfinished sets, H differs from it only by the "
              "Pudding standing, the mean-opponent term and its bounded scale, which were not separated; the card "
              "values act through the rollout policy B.")
    s += (f" S and I together are worth 10.2 ± 5.7 pp (<i>z</i> = 3.5). S alone, decoupled selection with v3's "
          f"determinisation, beats v3 by {fpp(SALONE)}, and I adds a further 3.8 ± 5.7 pp, a small gain, as Long et "
          "al.'s analysis {long} suggests for a game that reveals its hidden information within <i>n</i> − 1 turns. P "
          "and M cannot be removed, but a 2<sup>4</sup> factorial on v3 (R; M; <i>K</i> = 0.5 vs 1.41; completion "
          "horizon <i>T</i> = 8 vs 4; 16 × 500 games) gave main effects of +8.7 pp for R (<i>z</i> = 8.1), +5.2 for "
          "<i>K</i>, +4.1 for M and +3.6 for <i>T</i> = 8 (each ±2.1); of the six interactions only R × <i>T</i> "
          "(−4.2: with R, <i>T</i> = 8 no longer helps) and M × <i>K</i> (+3.2) exceed |<i>z</i>| = 2.")
    return s


def text_last_study():
    rest = [LAST[k] for k in ("X3", "N05", "V2D", "W1")]
    s = ("Of five later candidate changes (FIGREF_EFFECTS, bottom), four changed the win rate by at most "
         f"{max(abs(d['dw']) for d in rest):.1f} pp (each ±{max(d['dwci'] for d in rest):.1f}), among them card "
         "constants tuned for three players (X3), a stronger greedy policy (45.0% of greedy games against two X0 "
         "players); rollouts to "
         f"the game end with a win-gap value cost {fpp(LAST['R3V2'], sign=False)}, the only change that survives a "
         "Holm correction. As X3 also changes the Pudding weight of the leaf evaluation, its null result only weakly "
         "echoes simulation balancing {silver}.")
    return s


def text_budget():
    d = DEVD
    s = (f"In development runs (FIGREF_BUDGET; independent samples) v3 doubled v2's win rate at 100 ms "
         f"({DEV['v2'][0]['win']:.1f}% → {DEV['v3'][0]['win']:.1f}%) but lost {fpp(d['v3drop'], z=False, sign=False)} "
         f"from 100 to 900 ms, the final agent {fpp(d['v4drop'], z=False, sign=False)}; the gain from S and I rose "
         f"from {fpp(d['si100'], z=False)} to {fpp(d['si900'], z=False)}, by {fpp(d['rise'])}. We hypothesise that "
         "v3's decline comes from its sequential tree, in which later players react to cards chosen earlier in the "
         f"same turn: a wrong model that more iterations exploit more. S alone gained {fpp(d['s100'], z=False)} at "
         f"100 ms but {fpp(d['s500'], z=False)} at 500 ms")
    if (500, "v3") in ITER and (500, "v4") in ITER:
        s += (f", and v3 and the final agent ran similar numbers of iterations ({num(ITER[(500, 'v3')])} and "
              f"{num(ITER[(500, 'v4')])} per decision at 500 ms)")
    s += (f". As v3's {DEV['v3'][0]['win']:.1f}% at 100 ms was the best of the 16 factorial cells (45.1–49.4% on "
          "three new seed sets), the development rise is probably overstated.")
    s += (f" On seed block 7 the final agent won {OLD['3p100']['win']:.1f}%, {OLD['3p500']['win']:.1f}% and "
          f"{OLD['3p900']['win']:.1f}% at 100, 500 and 900 ms: its development decline did not replicate.")
    if GAIN["3p100"] and GAIN["3p500"] and RISE:
        z = RISE["zw"]
        verdict = ("which supports H2" if z >= 1.96 else "which contradicts H2" if z <= -1.96 else
                   "in the direction of H2 but not significant" if RISE["dw"] > 0 else
                   "so H2 is not supported on these deals")
        s += (f" On the held-out deals, played by both agents at each budget, v3 won {fwin(V3H['3p100'])} and "
              f"{fwin(V3H['3p500'])} at 100 and 500 ms, so S and I changed v3's win rate by "
              f"{fpp(GAIN['3p100'], z=False)} and {fpp(GAIN['3p500'], z=False)}: a difference of {fpp(RISE)}, "
              f"{verdict}")
        if abs(z) < 1.96:   # 80% power at alpha = 0.05: (1.96 + 0.84) SE
            s += f" (this test detects only changes of about {1.43 * RISE['dwci']:.0f} pp or more)"
        s += "."
    elif GAIN["3p100"]:
        s += f" On the held-out deals S and I gain {fpp(GAIN['3p100'])} at 100 ms (v3: {fwin(V3H['3p100'])})."
    ks = [k for k in ("3p500", "3p900") if BUD[k]]
    if ks:
        sig = any(abs(BUD[k]["zw"]) >= 1.96 for k in ks)
        parts = " and ".join(f"{fpp(BUD[k], z=False)} at {k[2:]} ms" for k in ks)
        dms = [BUD[k]["dm"] for k in ks]
        msig = any(abs(BUD[k]["zm"]) >= 1.96 for k in ks)
        if msig and (all(x < 0 for x in dms) or all(x > 0 for x in dms)):
            marg = (("and" if sig else "but") + " its score margin " + ("fell" if dms[0] < 0 else "rose") + " by "
                    + " and ".join(f"{abs(BUD[k]['dm']):.2f} ± {BUD[k]['dmci']:.2f}" for k in ks) + " points")
        else:
            marg = ("and its score margin changed by "
                    + " and ".join(f"{BUD[k]['dm']:+.2f} ± {BUD[k]['dmci']:.2f}" for k in ks).replace("-", "−")
                    + " points")
        s += (" The final agent's own win rate on the held-out deals "
              + ("changed" if sig else "did not change significantly") + f" with the budget ({parts} against 100 ms, "
              f"same deals), {marg}.")
    i3, i4 = ISM["v3"], ISM["v4"]
    if all(x is not None for x in i3 + i4):
        s += (f" Much of the fall of the curves comes from the opponent: TAG's ISMCTS won {i4[0]:.1f}% of its seats "
              f"against the final agent at 100 ms but {i4[2]:.1f}% at 900 ms (against v3: {i3[0]:.1f}% → "
              f"{i3[2]:.1f}%).")
    return s


def text_four():
    f = FOUR
    s = (f"In four-player development no change gave a significant gain at its own budget: X0 added {fpp(f['x0'], z=False)} at 100 ms (two runs of each pooled), X0 with W2 "
         f"and <i>K</i> = 1 added {fpp(f['w2'], z=False)} at 500 ms (two runs pooled) and F + L a further "
         f"{fpp(f['fl500'], z=False)}, whereas F + L cost {fpp(f['fl100'], z=False, sign=False)} at 100 ms; we kept "
         "each change where its estimate was positive. With three players at 500 ms (both arms with F + L), X0, W2 "
         f"and <i>K</i> = 1 cost {fpp(f['cost3'], sign=False)}, an effect {fpp(f['inter'], sign=False)} "
         "different from that with four. This supports H3, although the gains with four players were not "
         "significant; as the contrast was chosen after the four-player runs, it is exploratory. Depth matters – at 500 ms a tree cut to depth 1 lost "
         f"{fpp(f['depth1'], z=False, sign=False)} – and so does time up to 500 ms: with 100 ms (and its 100 ms "
         f"settings) our agent won {fpp(f['b100'], z=False, sign=False)} less (<i>z</i> = {abs(f['b100']['zw']):.1f}), "
         "but 1 000 ms (with <i>K</i> = 1.37) gave "
         f"{fpp(f['t1000'], z=False)}.")
    if DIAG:
        s += (f" On {DIAG['states']} positions from random games the 500 ms tree has {DIAG['kids']} joint-move "
              f"children at the root on average, {DIAG['once']}% of them visited once (mean depth "
              f"{DIAG['depth']:.1f}); doubling the time or the rollouts per leaf changed the card in "
              f"{count73(DIAG['t1000'])} and {count73(DIAG['L16'])} of them, no more often than a re-run with the "
              f"same seed ({count73(DIAG['rerun'])}) or another seed ({count73(DIAG['seedSame'])}).")
    return s


def text_fig_budget_caption():
    return ("Win rate with three players against v2 and TAG's ISMCTS by CPU time per decision. Lines: development runs "
            "(seed 1, independent samples; 250 games at 900 ms). Open markers: the final agent and v3 on the held-out "
            "deals (seed block 100 000). Bars: Wilson 95% intervals.")


def text_fig_effects_caption():
    return ("Paired effects with three players at 100 ms (seed block 7, 500 games each), ± 95% CI; filled markers "
            "survive a Holm correction within their group. Top: one component removed from the final agent (base: "
            f"final agent, {OLD['3p100']['win']:.1f}%). Bottom: TAG's lead heuristic and five later candidate changes, "
            f"against a re-run of the final agent on the machine of the held-out runs ({BASE5['win']:.1f}%).")


def text_conclusion_1():
    s = ("For RQ1, removing any component lowers the final agent's score margin significantly, but after a Holm "
         "correction only removing the evaluation H (worth "
         f"{abs(LEAD['dw']):.1f} pp over TAG's lead heuristic) or the decoupled information-set tree (10.2 pp over "
         "v3's sequential, determinised tree) also lowers its win rate. On held-out deals the final agent wins "
         f"{fwin(VAL['3p100'])}, {fwin(VAL['3p500'])} and {fwin(VAL['3p900'])} of three-player games at 100, 500 and "
         f"900 ms against TAG's ISMCTS and our previous agent, and {fwin(VAL['4p100'])} of four-player games at "
         "100 ms.")
    if RISE and GAIN["3p100"] and GAIN["3p500"]:
        g = f"from {fpp(GAIN['3p100'], z=False)} at 100 ms to {fpp(GAIN['3p500'], z=False)} at 500 ms"
        z = RISE["zw"]
        if z >= 1.96:
            s += f" For RQ2, the gain from S and I grew {g} on the held-out deals (<i>z</i> = {z:.1f})."
        elif z <= -1.96:
            s += (f" For RQ2, the gain from S and I shrank {g} on the held-out deals (<i>z</i> = {abs(z):.1f}), "
                  "contrary to H2.")
        elif RISE["dw"] > 0:
            s += (f" For RQ2, the gain from S and I grew {g} on the held-out deals, as in development but not "
                  f"significantly (<i>z</i> = {z:.1f}).")
        else:
            s += (f" For RQ2, the gain from S and I did not grow on the held-out deals ({g}), although in development "
                  "v3 had lost more than the final agent from 100 to 900 ms (a difference that was not significant).")
    else:
        s += (" For RQ2, in development v3's sequential tree lost more strength at longer budgets than the decoupled "
              "one, though the difference between the trends is not significant.")
    s += (f" For RQ3, settings tuned for four players cost {fpp(FOUR['cost3'], z=False, sign=False)} with three "
          "players while their gain with four was not significant (an exploratory result), and with four players time "
          "beyond 500 ms rarely changes the decision of the wide, shallow joint-move tree.")
    return s.replace("= -", "= −")


def text_conclusion_2():
    return ("Our opponents were framework agents and our own previous version, the four-player settings rest on "
            "250-game runs and non-significant gains, and CPU-time budgets hide wall-clock effects. As hand-tuned "
            "changes to the model no longer move the agent, future work should learn the evaluation from game "
            "outcomes, e.g. by regressing final ranks on features of round-end states; regret matching or EXP3 "
            "deserve a test, as decoupled UCT is best overall but not in every game {tak} and was the most "
            "exploitable in Goofspiel {lanctot}; and opponent models fitted to observed picks could replace the "
            "single greedy policy that now models all players.")


PLACEHOLDERS = {
    "FINAL_AGENT_PARAGRAPH": text_final_agent, "ABLATION_PARAGRAPH": text_ablation, "BUDGET_PARAGRAPH": text_budget,
    "FOUR_PARAGRAPH": text_four, "LAST_STUDY_PARAGRAPH": text_last_study,
    "FIG1_CAPTION": text_fig1_caption, "FIG_BUDGET_CAPTION": text_fig_budget_caption,
    "FIG_EFFECTS_CAPTION": text_fig_effects_caption, "CONCLUSION_1": text_conclusion_1,
    "CONCLUSION_2": text_conclusion_2,
}


# ----------------------------------------------------------------------------------------------- helpers
def number_citations(blocks):
    order = []

    def walk(text):
        for k in re.findall(r"\{(\w+)\}", text):
            if k not in REFS:
                sys.exit(f"unknown citation {k}")
            if k not in order:
                order.append(k)

    for b in blocks:
        if b[0] in ("p", "bullet", "pai"):
            walk(b[1])
        elif b[0] == "table":
            for row in b[4]:
                for c in row:
                    walk(c)
            walk(b[1])
        elif b[0] == "fig":
            walk(b[3])
    num = {k: i + 1 for i, k in enumerate(order)}
    unused = [k for k in REFS if k not in num]
    if unused:
        sys.exit(f"references never cited: {unused}")
    return order, num


def cite(text, num):
    text = re.sub(r"\{(\w+)\}, \{(\w+)\}", lambda m: f"[{num[m.group(1)]}], [{num[m.group(2)]}]", text)
    return re.sub(r"\{(\w+)\}", lambda m: f"[{num[m.group(1)]}]", text)


TOKEN = re.compile(r"(<i>|</i>|<b>|</b>|<sub>|</sub>|<sup>|</sup>)")


def add_runs(par, text, size=None, font=None, bold=None, italic=None):
    """Adds text with <i>, <b>, <sub>, <sup> markup as runs of paragraph par."""
    st = {"i": False, "b": False, "sub": False, "sup": False}
    for tok in TOKEN.split(text):
        if not tok:
            continue
        m = re.fullmatch(r"<(/?)(i|b|sub|sup)>", tok)
        if m:
            st[m.group(2)] = m.group(1) == ""
            continue
        r = par.add_run(tok)
        if st["i"] or italic:
            r.italic = True
        if st["b"] or bold:
            r.bold = True
        if st["sub"]:
            r.font.subscript = True
        if st["sup"]:
            r.font.superscript = True
        if size:
            r.font.size = Pt(size)
        if font:
            r.font.name = font
            r._element.rPr.rFonts.set(qn("w:cs"), font)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return par


def sty(doc, name):
    """Style by its displayed name (LibreOffice writes 'Heading 1', which python-docx's lookup does not find)."""
    for s in doc.styles:
        if s.name == name:
            return s
    raise KeyError(name)


def new_par(doc, anchor, style):
    par = doc.add_paragraph()
    par.style = sty(doc, style)
    anchor.addprevious(par._p)
    return par


def set_tabs(par, stops):
    ppr = par._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for kind, pos in [("clear", 720), ("clear", 2520), ("clear", 5040)] + stops:
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), kind)
        t.set(qn("w:pos"), str(pos))
        tabs.append(t)
    ppr.append(tabs)


def border(el, tag, val="single", sz=6):
    b = OxmlElement(f"w:{tag}")
    b.set(qn("w:val"), val)
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "0")
    b.set(qn("w:color"), "000000")
    el.append(b)


def add_table(doc, anchor, caption, widths, header, rows, note, num):
    cap = new_par(doc, anchor, "table head")
    cap.add_run(caption)
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=len(rows) + 1, cols=len(header))
    anchor.addprevious(t._tbl)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    border(borders, "top", sz=8)
    border(borders, "bottom", sz=8)
    tblPr.append(borders)
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for ci, w in enumerate(widths):
        for row in t.rows:
            row.cells[ci].width = Inches(w)
    for ri, values in enumerate([header] + rows):
        for ci, v in enumerate(values):
            cell = t.rows[ri].cells[ci]
            par = cell.paragraphs[0]
            par.style = sty(doc, "table col head" if ri == 0 else "table copy")
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci <= 1 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(par, cite(v, num))
            tcPr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side in ("left", "right"):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:w"), "50")
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            tcPr.append(mar)
            if ri == 0:
                tb = OxmlElement("w:tcBorders")
                border(tb, "bottom", sz=6)
                tcPr.append(tb)
    for row in t.rows[:-1]:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
    if note:
        for cell in t.rows[-1].cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True
        np_ = new_par(doc, anchor, "table copy")
        np_.alignment = WD_ALIGN_PARAGRAPH.LEFT
        np_.paragraph_format.space_before = Pt(2)
        np_.paragraph_format.space_after = Pt(8)
        add_runs(np_, note)
    else:
        sp = new_par(doc, anchor, "table copy")
        sp.paragraph_format.space_after = Pt(4)


def fill_placeholders():
    rep = {k: f() for k, f in PLACEHOLDERS.items()}
    figs = [b[1] for b in B if b[0] == "fig"]
    for i, f in enumerate(figs):
        for key in ("turn", "effects", "budget"):
            if key in f:
                rep["FIGREF_" + key.upper()] = f"Fig. {i + 1}"

    def sub(t):
        for _ in range(2):  # placeholders may contain FIGREF tokens
            for k, val in rep.items():
                t = t.replace(k, val)
        t = t.replace("<i>Sushi Go!</i>", "Sushi Go!").replace("Sushi Go!", "<i>Sushi Go!</i>")  # game name in italics
        return re.sub(r"  +", " ", t).replace(" .", ".").strip()

    for i, b in enumerate(B):
        if b[0] in ("p", "pai", "bullet"):
            B[i] = (b[0], sub(b[1])) + tuple(b[2:])
        elif b[0] == "fig":
            B[i] = (b[0], b[1], b[2], sub(b[3]))
        elif b[0] == "table":
            B[i] = (b[0], b[1], b[2], b[3], [[sub(c) for c in row] for row in b[4]], b[5])


def build():
    fill_placeholders()
    os.makedirs(WORK, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "docx:MS Word 2007 XML", "--outdir", WORK, TEMPLATE],
                   check=True, capture_output=True)
    doc = docx.Document(os.path.join(WORK, "Paper_Template.docx"))
    order, num = number_citations(B)
    P = doc.paragraphs

    # title
    for r in P[0].runs[1:]:
        r._r.getparent().remove(r._r)
    P[0].runs[0].text = TITLE.split("Card-Value")[0]
    P[0].runs[0].font.name = "Times New Roman"
    P[0].runs[0]._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    P[0].runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if "Card-Value" in TITLE:
        tail = copy.deepcopy(P[0].runs[0]._r)
        for t in tail.findall(qn("w:t")):
            tail.remove(t)
        t1 = OxmlElement("w:t")
        t1.text = "Card"
        tail.append(t1)
        tail.append(OxmlElement("w:noBreakHyphen"))
        t2 = OxmlElement("w:t")
        t2.set(qn("xml:space"), "preserve")
        t2.text = "Value" + TITLE.split("Card-Value")[1]
        tail.append(t2)
        P[0].runs[0]._r.addnext(tail)
    # subtitle note and empty author line of section 0
    for i in (1, 2):
        P[i]._p.getparent().remove(P[i]._p)
    # authors: three columns
    a_par = P[5]
    for i in (4, 6):
        P[i]._p.getparent().remove(P[i]._p)
    # one paragraph with column breaks between the authors, so that all three blocks start at the same height
    for r in list(a_par._p.findall(qn("w:r"))):
        a_par._p.remove(r)
    for k, lines in enumerate(AUTHORS):
        if k > 0:
            a_par.add_run().add_break(docx.enum.text.WD_BREAK.COLUMN)
        for j, line in enumerate(lines):
            r = a_par.add_run(line)
            r.font.size = Pt(9)
            r.italic = j >= 2
            if j < len(lines) - 1:
                r.add_break()
    # empty paragraphs of section 2
    for i in (8, 9):
        P[i]._p.getparent().remove(P[i]._p)
    # template body: everything between the abstract and the section break of the 2-column section
    sect_par = P[85]._p
    body = doc.element.body
    el = P[11]._p
    while el is not sect_par:
        nxt = el.getnext()
        body.remove(el)
        el = nxt
    anchor = sect_par
    # first-page footer: the IEEE copyright line of the template becomes our note on the use of generative AI
    fpars = doc.sections[0].first_page_footer.paragraphs
    for par in fpars:
        for r in par.runs:
            r.text = ""
    note = cite(FIRST_PAGE_NOTE, num)
    fp = fpars[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_runs(fp, note, size=8)
    # top margins as in the module template: 27 pt above the title, 54 pt on the following (body) pages
    for k, sec in enumerate(doc.sections):
        sec.top_margin = Pt(27 if k == 0 else 54)

    ab = new_par(doc, anchor, "Abstract")
    r = ab.add_run("Abstract—")
    r.italic = True
    add_runs(ab, ABSTRACT)
    kw = new_par(doc, anchor, "Keywords")
    kw.add_run("Keywords—")
    r = kw.add_run(KEYWORDS)
    r.bold = False

    for b in B:
        kind = b[0]
        if kind == "h1":
            new_par(doc, anchor, "Heading 1").add_run(b[1])
        elif kind == "h2":
            new_par(doc, anchor, "Heading 2").add_run(b[1])
        elif kind == "h5":
            hp = new_par(doc, anchor, "Heading 5")
            hp.paragraph_format.keep_with_next = True
            hp.add_run(b[1])
        elif kind == "p":
            add_runs(new_par(doc, anchor, "Body Text"), cite(b[1], num))
        elif kind == "pai":
            add_runs(new_par(doc, anchor, "Body Text"), cite(b[1], num))
        elif kind == "bullet":
            par = new_par(doc, anchor, "bullet list")
            ppr = par._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            for tag, val in (("w:ilvl", "0"), ("w:numId", "2")):
                e = OxmlElement(tag)
                e.set(qn("w:val"), val)
                numpr.append(e)
            ppr.append(numpr)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "576")
            ind.set(qn("w:hanging"), "288")
            ppr.append(ind)
            add_runs(par, cite(b[1], num))
        elif kind == "eq":
            par = new_par(doc, anchor, "equation")
            set_tabs(par, [("center", 2438), ("right", 4877)])
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(4)
            par.add_run("\t")
            add_runs(par, b[1], font="Times New Roman")
            r = par.add_run(f"\t({b[2]})")
            r.font.name = "Times New Roman"
        elif kind == "table":
            add_table(doc, anchor, b[1], b[2], b[3], b[4], b[5], num)
        elif kind == "fig":
            par = new_par(doc, anchor, "Normal")
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(6)
            par.paragraph_format.keep_with_next = True
            par.add_run().add_picture(os.path.join(HERE, b[1]), width=Inches(b[2]))
            cp = new_par(doc, anchor, "figure caption")
            cp.paragraph_format.keep_together = True
            add_runs(cp, cite(b[3], num))
        elif kind == "refs":
            for k in order:
                add_runs(new_par(doc, anchor, "references"), REFS[k])

    doc.core_properties.title = TITLE
    doc.core_properties.author = "ECS7032P Group A"
    doc.save(OUT + ".docx")
    pdf_filter = ('pdf:writer_pdf_Export:{"UseLosslessCompression":{"type":"boolean","value":"true"},'
                  '"ReduceImageResolution":{"type":"boolean","value":"false"}}')
    subprocess.run(["soffice", "--headless", "--convert-to", pdf_filter, "--outdir", HERE, OUT + ".docx"],
                   check=True, capture_output=True)
    print("written", OUT + ".docx", OUT + ".pdf")
    words = lambda t: len(re.sub(r"<[^>]+>|\{\w+\}", "", t).split())
    print("abstract words:", words(ABSTRACT))
    sec, counts = None, {}
    for b in B:
        if b[0] in ("h1", "h5"):
            sec = b[1]
        elif b[0] in ("p", "bullet") and sec:
            counts[sec] = counts.get(sec, 0) + words(b[1])
    print("words per section:", counts)


if __name__ == "__main__":
    build()
